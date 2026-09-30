from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib import colors

from utils.document_utils import sanitize_text

ROOT = Path(__file__).resolve().parents[1]
LOGO = ROOT / "assets" / "logo.svg"


def _title_from_text(text: str, fallback: str) -> str:
    first = next((line.strip() for line in sanitize_text(text).splitlines() if line.strip()), "")
    return first[:120] if first else fallback


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    if LOGO.with_suffix(".png").exists():
        logo_p = document.add_paragraph()
        logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_p.add_run().add_picture(str(LOGO.with_suffix(".png")), width=Inches(4.0))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(_title_from_text(text, doc_type).upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("LegalEase AI - Draft Document")
    sr.italic = True
    sr.font.size = Pt(9)

    for block in re.split(r"\n\s*\n", sanitize_text(text)):
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        if len(lines) == 1 and (
            lines[0].isupper() or re.match(r"^\d+(?:\.\d+)*\.?\s+", lines[0])
        ):
            p = document.add_paragraph()
            r = p.add_run(lines[0])
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
        else:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.line_spacing = 1.15
            for i, line in enumerate(lines):
                if i:
                    p.add_run().add_break()
                p.add_run(line)

    clean_terms = [x.strip() for x in terms.split(";") if x.strip()]
    if clean_terms:
        document.add_paragraph()
        h = document.add_paragraph()
        rr = h.add_run("KEY TERMS")
        rr.bold = True
        table = document.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"
        for idx, term in enumerate(clean_terms, start=1):
            cells = table.add_row().cells
            cells[0].text = str(idx)
            cells[1].text = term

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(
        "LegalEase AI - AI-assisted draft. Review by a qualified legal professional is recommended."
    )
    fr.font.size = Pt(8)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


class NumberedCanvasMixin:
    pass


def _pdf_footer(canvas, doc):
    canvas.saveState()
    png_logo = LOGO.with_suffix(".png")
    if png_logo.exists():
        canvas.drawImage(str(png_logo), 15 * mm, A4[1] - 17 * mm, width=42 * mm, height=11 * mm, preserveAspectRatio=True, mask="auto")
    canvas.setFont("Helvetica", 8)
    canvas.drawCentredString(
        A4[0] / 2,
        10 * mm,
        "LegalEase AI - AI-assisted draft - Legal review recommended",
    )
    canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"Page {doc.page}")
    canvas.restoreState()


def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    output = BytesIO()
    doc = BaseDocTemplate(
        output,
        pagesize=A4,
        rightMargin=17 * mm,
        leftMargin=17 * mm,
        topMargin=20 * mm,
        bottomMargin=18 * mm,
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates([PageTemplate(id="legal", frames=frame, onPage=_pdf_footer)])

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "LegalTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        alignment=TA_CENTER,
        spaceAfter=8,
    )
    heading = ParagraphStyle(
        "LegalHeading",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        spaceBefore=8,
        spaceAfter=5,
    )
    body = ParagraphStyle(
        "LegalBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        spaceAfter=6,
    )
    story = [
        Paragraph(_title_from_text(text, doc_type).upper(), title_style),
        Paragraph("LegalEase AI - Draft Document", body),
        Spacer(1, 4),
    ]

    for block in re.split(r"\n\s*\n", sanitize_text(text)):
        block = block.strip()
        if not block:
            continue
        escaped = block.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        lines = escaped.splitlines()
        if len(lines) == 1 and (
            lines[0].isupper() or re.match(r"^\d+(?:\.\d+)*\.?\s+", lines[0])
        ):
            story.append(Paragraph(lines[0], heading))
        else:
            story.append(Paragraph("<br/>".join(lines), body))

    clean_terms = [x.strip() for x in terms.split(";") if x.strip()]
    if clean_terms:
        data = [["No.", "Key term"]]
        data.extend([[str(i), t] for i, t in enumerate(clean_terms, 1)])
        table = Table(data, colWidths=[15 * mm, 150 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.extend([Spacer(1, 8), Paragraph("KEY TERMS", heading), table])

    doc.build(story)
    return output.getvalue()
